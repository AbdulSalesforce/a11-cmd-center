#!/usr/bin/env python3
"""
ACR Generator — audit tool edition.
Reads findings from a JSON file written by the Node server, then populates
the ACR Word template using the shared boilerplate CSV.

Usage (called by the server):
  python generate_acr.py <findings_json> <output_docx> <product_name>
"""
import sys
import os
import json
import csv
import re
from docx import Document
from collections import defaultdict
from datetime import date

# Paths to shared assets — same files used by the desktop script
ASSETS_DIR = os.path.join(
    os.path.expanduser("~"), "Desktop", "ACRGeneration"
)
TEMPLATE_PATH = os.path.join(ASSETS_DIR, "[Template v2026-Apr-20] ACR.docx")
BOILERPLATE_CSV = os.path.join(ASSETS_DIR, "Copy of ACR Master Library - WCAG Table Criteria.csv")


def load_findings_from_json(json_path):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    # data is a list of {page, wcag, description}
    print(f"Loaded {len(data)} findings from {json_path}")
    return data


def load_boilerplate():
    boilerplate = {}
    section508_ch3 = {}
    section508_ch6 = {}
    en301 = {}

    with open(BOILERPLATE_CSV, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            table_name  = row.get("Table", "").strip()
            criteria_id = row.get("Criteria ID", "").strip()
            conformance = row.get("Conformance Level", "").strip()
            remark      = row.get("Standard Remark (Boilerplate)", "").strip()
            mapped_raw  = row.get("Mapped WCAG SC", "").strip()

            wcag_match = re.search(r"(\d+\.\d+\.\d+)", criteria_id)

            if wcag_match and ("Table 1" in table_name or "Table 2" in table_name):
                wcag_num = wcag_match.group(1)
                boilerplate[wcag_num] = {"conformance": conformance, "remark": remark, "table": table_name}

            if "Section 508 Chapter 3" in table_name:
                m = re.match(r"^(\d{3}(?:\.\d+)?)", criteria_id)
                if m:
                    mapped = [w.strip() for w in mapped_raw.split(",") if w.strip()]
                    section508_ch3[m.group(1)] = {"conformance": conformance, "remark": remark, "mapped_wcag": mapped}

            if "Section 508 Chapter 6" in table_name:
                m = re.match(r"^(\d{3}\.\d+)", criteria_id)
                if m:
                    section508_ch6[m.group(1)] = {"conformance": conformance, "remark": remark}

            if "EN 301 549 Chapter 4" in table_name:
                m = re.match(r"^(\d+\.\d+\.\d+(?:\.\d+)?)", criteria_id)
                if m:
                    mapped = [w.strip() for w in mapped_raw.split(",") if w.strip()]
                    en301[m.group(1)] = {"conformance": conformance, "remark": remark, "mapped_wcag": mapped}

    print(f"Boilerplate loaded — WCAG: {len(boilerplate)}, 508-Ch3: {len(section508_ch3)}, "
          f"508-Ch6: {len(section508_ch6)}, EN301: {len(en301)}")
    return boilerplate, section508_ch3, section508_ch6, en301


def extract_wcag_number(text):
    m = re.search(r"(\d+\.\d+\.\d+)", text)
    return m.group(1) if m else None


def extract_508_or_en301_id(text):
    m = re.match(r"^(\d{3}(?:\.\d+)?)", text)
    if m:
        return m.group(1)
    m = re.match(r"^(\d+\.\d+\.\d+(?:\.\d+)?)", text)
    return m.group(1) if m else None


def create_summary(page, description):
    desc = description.strip()
    for marker in ["Remediation:", "Steps to Reproduce:", "Impact:", "Recommendations:"]:
        if marker in desc:
            desc = desc.split(marker)[0].strip()
            break
    sentences = re.split(r"(?<=[.!?])\s+(?=[A-Z])", desc)
    summary = ""
    for s in sentences:
        s = s.strip()
        if s and len(s) > 15 and not s.startswith(("-", "•", "*")):
            summary = s
            break
    if not summary:
        summary = desc.split("\n")[0].strip() if "\n" in desc else desc[:200].strip()
    if len(summary) > 250:
        truncated = summary[:250]
        last_period = truncated.rfind(".")
        summary = summary[:last_period + 1] if last_period > 100 else truncated.rsplit(" ", 1)[0].strip() + "."
    if not summary.endswith((".", "!", "?")):
        summary = summary.rstrip(",;:") + "."
    return f"{page}: {summary}"


def group_findings_by_wcag(findings):
    grouped = defaultdict(list)
    for f in findings:
        wcag_num = extract_wcag_number(f["wcag"])
        if wcag_num:
            grouped[wcag_num].append(create_summary(f["page"], f["description"]))
    return grouped


def replace_product(text, product_name):
    if not text or product_name == "[Product Name]":
        return text
    for placeholder in ("[Product Name]", "[Product]", "[product]"):
        text = text.replace(placeholder, product_name)
    return text


def set_font(paragraph):
    from docx.shared import Pt
    for run in paragraph.runs:
        run.font.name = "Salesforce Sans"
        run.font.size = Pt(11)


def populate_acr(grouped_findings, boilerplate, s508_ch3, s508_ch6, en301,
                 output_path, product_name):
    doc = Document(TEMPLATE_PATH)

    for table in doc.tables:
        if len(table.columns) != 3:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if "Criteria" not in headers[0]:
            continue

        first = table.rows[1].cells[0].text.strip() if len(table.rows) > 1 else ""
        is_wcag     = bool(re.search(r"\d+\.\d+\.\d+", first))
        is_508_ch3  = "Section 508 Chapter 3" in table.rows[0].cells[0].text or first.startswith("302.")
        is_508_ch6  = ("Section 508 Chapter 6" in table.rows[0].cells[0].text or
                       any(first.startswith(p) for p in ("601.", "602.", "603.")))
        is_en301    = ("EN 301 549 Chapter 4" in table.rows[0].cells[0].text or
                       (first.startswith("4.2.") and "Usage" in first))

        for row in table.rows[1:]:
            conf_cell    = row.cells[1]
            remarks_cell = row.cells[2]
            criteria_txt = row.cells[0].text.strip()

            if is_wcag and not is_508_ch3 and not is_508_ch6 and not is_en301:
                wcag_num = extract_wcag_number(criteria_txt)
                if not wcag_num:
                    continue
                if wcag_num in grouped_findings:
                    conf_cell.text = "Partially Supports"
                    set_font(conf_cell.paragraphs[0])
                    remarks_cell.text = ""
                    if wcag_num in boilerplate:
                        p = remarks_cell.paragraphs[0]
                        p.text = replace_product(boilerplate[wcag_num]["remark"], product_name)
                        set_font(p)
                        ep = remarks_cell.add_paragraph()
                        run = ep.add_run(f"However, in {product_name}, there are exceptions:")
                        run.bold = True
                        set_font(ep)
                    else:
                        p = remarks_cell.paragraphs[0]
                        run = p.add_run(f"However, in {product_name}, there are exceptions:")
                        run.bold = True
                        set_font(p)
                    for summary in grouped_findings[wcag_num]:
                        fp = remarks_cell.add_paragraph()
                        fp.text = f"• {summary}"
                        set_font(fp)
                elif wcag_num in boilerplate:
                    bp = boilerplate[wcag_num]
                    conf_cell.text = bp["conformance"]
                    set_font(conf_cell.paragraphs[0])
                    remarks_cell.text = replace_product(bp["remark"], product_name)
                    set_font(remarks_cell.paragraphs[0])

            elif is_508_ch3:
                s508_id = extract_508_or_en301_id(criteria_txt)
                if not s508_id or s508_id not in s508_ch3:
                    continue
                bp = s508_ch3[s508_id]
                has_findings = any(w in grouped_findings for w in bp.get("mapped_wcag", []))
                if has_findings:
                    conf_cell.text = "Partially Supports"
                    set_font(conf_cell.paragraphs[0])
                    remarks_cell.text = "See information in WCAG 2.2 section."
                    set_font(remarks_cell.paragraphs[0])
                else:
                    conf_cell.text = bp["conformance"]
                    set_font(conf_cell.paragraphs[0])
                    remarks_cell.text = replace_product(bp["remark"], product_name)
                    set_font(remarks_cell.paragraphs[0])

            elif is_508_ch6:
                s508_id = extract_508_or_en301_id(criteria_txt)
                if not s508_id or s508_id not in s508_ch6:
                    continue
                bp = s508_ch6[s508_id]
                conf_cell.text = bp["conformance"]
                set_font(conf_cell.paragraphs[0])
                remarks_cell.text = replace_product(bp["remark"], product_name)
                set_font(remarks_cell.paragraphs[0])

            elif is_en301:
                en_id = extract_508_or_en301_id(criteria_txt)
                if not en_id or en_id not in en301:
                    continue
                bp = en301[en_id]
                has_findings = any(w in grouped_findings for w in bp.get("mapped_wcag", []))
                if has_findings:
                    conf_cell.text = "Partially Supports"
                    set_font(conf_cell.paragraphs[0])
                    remarks_cell.text = "See information in WCAG 2.2 section."
                    set_font(remarks_cell.paragraphs[0])
                else:
                    conf_cell.text = bp["conformance"]
                    set_font(conf_cell.paragraphs[0])
                    remarks_cell.text = replace_product(bp["remark"], product_name)
                    set_font(remarks_cell.paragraphs[0])

    doc.save(output_path)
    print(f"ACR saved to {output_path}")


def main():
    if len(sys.argv) < 4:
        print("Usage: generate_acr.py <findings.json> <output.docx> <product_name>")
        sys.exit(1)

    findings_json = sys.argv[1]
    output_path   = sys.argv[2]
    product_name  = sys.argv[3]

    findings         = load_findings_from_json(findings_json)
    bp, s508c3, s508c6, en301 = load_boilerplate()
    grouped          = group_findings_by_wcag(findings)

    populate_acr(grouped, bp, s508c3, s508c6, en301, output_path, product_name)


if __name__ == "__main__":
    main()
